"""Tests for deep candidate extraction, preview bridge, and semantic pipeline helpers."""
import io
import zipfile
from unittest.mock import AsyncMock, patch

import pytest
from docx import Document

from app.deid.discovery.deep_candidates import extract_deep_candidates
from app.deid.discovery.deep_flows import (
    _assign_risk_ids,
    _reject_cross_paragraph,
    run_deep_suggest_all,
)
from app.deid.discovery.semantic_rules import extract_program_risks
from app.deid.engine.pipeline import partition_semantic_pairs, run_deid_pipeline
from app.deid.engine.preview import assign_placeholder_map, build_preview_text
from app.deid.engine.pipeline import extract_sample_text
from app.models_deid import DeidJob, DeidJobEntity


def _minimal_docx(path, paragraphs: list[str]) -> None:
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    doc.save(path)


def test_extract_windows_finds_project_line():
    text = "项目名称：251231内控审计\n金额单位：万元"
    windows = extract_deep_candidates(text, window_size=400)
    assert any("251231" in w["text"] for w in windows)


def test_extract_windows_finds_stock_code():
    text = "持有证券600642.SHA股\n"
    windows = extract_deep_candidates(text, window_size=400)
    assert windows


def test_extract_windows_skips_plain_amount_line():
    text = "金额单位：万元\n营业收入：1000"
    windows = extract_deep_candidates(text, window_size=400)
    assert windows == []


def test_assign_risk_ids_dedupes():
    risks = _assign_risk_ids(
        [
            {"category": "project_id", "original": "251231内控审计", "note": "a"},
            {"category": "project_id", "original": "251231内控审计", "note": "b"},
        ]
    )
    assert len(risks) == 1
    assert risks[0]["risk_id"].startswith("r-")


def test_build_preview_text_replaces_entities():
    entities = [
        {
            "id": 1,
            "canonical_name": "中国能建",
            "entity_type": "company",
            "aliases": ["中国能建"],
            "is_excluded": False,
        },
    ]
    type_prefix_map = {"company": "公司"}
    sample = "中国能建审计报告"
    preview = build_preview_text(sample, entities, [], [], type_prefix_map=type_prefix_map)
    assert "中国能建" not in preview
    assert "[公司_1]" in preview


def test_placeholder_map_matches_assign(db):
    job = DeidJob(
        status="scanned",
        original_filename="t.docx",
        stored_path="deid/1/original_t.docx",
        pack_ids_json="[]",
    )
    db.add(job)
    db.flush()
    ents = [
        DeidJobEntity(
            job_id=job.id,
            canonical_name="乙公司",
            entity_type="company",
            source="manual",
            is_excluded=False,
        ),
        DeidJobEntity(
            job_id=job.id,
            canonical_name="甲公司",
            entity_type="company",
            source="manual",
            is_excluded=False,
        ),
    ]
    db.add_all(ents)
    db.commit()

    from app.deid.service import _assign_placeholders, _type_prefix_map

    _assign_placeholders(db, job.id)
    db.commit()
    rows = (
        db.query(DeidJobEntity)
        .filter(DeidJobEntity.job_id == job.id, DeidJobEntity.is_excluded.is_(False))
        .all()
    )
    db_map = {r.id: r.placeholder for r in rows}

    ent_dicts = [
        {"id": e.id, "entity_type": e.entity_type, "canonical_name": e.canonical_name}
        for e in sorted(ents, key=lambda x: x.id)
    ]
    preview_map = assign_placeholder_map(ent_dicts, type_prefix_map=_type_prefix_map(db))
    assert preview_map == db_map


def test_program_rules_disabled_by_default():
    risks = extract_program_risks("当前数据来源：同花顺")
    assert risks == []


def test_program_rules_tonghuashun_when_enabled(monkeypatch):
    monkeypatch.setenv("DEID_SEMANTIC_PROGRAM_RULES", "1")
    risks = extract_program_risks("当前数据来源：同花顺")
    assert any(
        r["category"] == "data_source" and "同花顺" in r["original"] for r in risks
    )


def test_default_semantic_rewrite_stock_code():
    from app.deid.discovery.semantic_rules import (
        apply_default_rewrites,
        default_semantic_rewrite,
        validate_suggest_rewrite,
    )

    assert default_semantic_rewrite("listing_code", "600182.SH") == "证券代码"
    assert default_semantic_rewrite("listing_structure", "流通A股,流通H股") == "多市场上市"
    risks = apply_default_rewrites(
        [{"category": "listing_code", "original": "833042.NQ", "source": "worker"}]
    )
    assert risks[0]["rewritten"] == "证券代码"
    assert not validate_suggest_rewrite("listing_code", "600182.SH", "S 佳通")
    assert validate_suggest_rewrite("data_source", "同花顺", "外部数据源")


def test_cross_paragraph_original_rejected():
    from app.deid.discovery.deep_flows import _reject_cross_paragraph

    assert _reject_cross_paragraph("第一行\n第二行")
    assert not _reject_cross_paragraph("单行文本")


def test_reject_low_quality_semantic_risks():
    from app.deid.discovery.deep_flows import _reject_low_quality_risk

    assert _reject_low_quality_risk("[公司_70](本期:2025-09-30)")
    assert _reject_low_quality_risk("公司名称")
    assert _reject_low_quality_risk("同花顺取数口径模糊的问题")
    assert _reject_low_quality_risk("导致退市风险警示的原因以及公司拟采取的应对措施")
    assert not _reject_low_quality_risk("600182.SH")
    assert not _reject_low_quality_risk("当前数据来源：同花顺")


@pytest.mark.asyncio
async def test_deep_detect_reuses_scan_chunks_not_per_line_windows():
    from app.deid.discovery.deep_flows import run_deep_detect
    from app.deid.discovery.flows import FlowResult
    from app.deid.discovery.llm import build_scan_chunk_plan, count_llm_chunks

    sample = "项目名称：251231内控审计\n" + "持有600642.SHA股\n" * 30

    async def fake_flow(*_args, **_kwargs):
        units = _args[1]
        assert len(units) == count_llm_chunks(sample)
        assert all(u.meta.get("kind") == "chunk" for u in units)
        return FlowResult(items=[], chunks=len(units))

    plan = build_scan_chunk_plan(sample)
    with patch("app.deid.discovery.deep_flows.run_worker_flow", new=fake_flow), patch(
        "app.deid.discovery.deep_flows.get_flow_prompt", return_value="test prompt"
    ), patch("app.deid.discovery.deep_flows.extract_program_risks", return_value=[]):
        risks, summary = await run_deep_detect(
            sample,
            object(),
            object(),
            job_id=1,
            scan_chunk_plan=plan,
        )
    assert risks == []
    assert summary.get("mode") == "scan_chunks_reuse"


@pytest.mark.asyncio
async def test_run_deep_suggest_all_fills_rewrites():
    risks = _assign_risk_ids(
        [{"category": "project_id", "original": "251231内控审计", "note": "a"}]
    )
    with patch(
        "app.deid.discovery.deep_flows.run_deep_suggest",
        new=AsyncMock(return_value="某审计项目"),
    ):
        out = await run_deep_suggest_all(
            risks,
            "sample",
            router=object(),
            db=object(),
            job_id=1,
        )
    assert out[0]["suggested_rewrite"] == "内控审计项目"
    assert out[0]["rewritten"] == "内控审计项目"


def test_dry_run_reports_missed(tmp_path):
    doc_path = tmp_path / "sem.docx"
    _minimal_docx(doc_path, ["项目名称：251231内控审计", "无关段落"])
    out_path = tmp_path / "out.docx"
    entities: list[dict] = []
    pairs = [
        {"original": "251231内控审计", "rewritten": "内控审计项目", "category": "project_id"},
        {"original": "不存在的内容", "rewritten": "替换", "category": "table_row"},
    ]
    result = run_deid_pipeline(doc_path, out_path, entities, [], [], semantic_pairs=pairs)
    assert result["semantic_applied_count"] == 1
    assert result["semantic_missed_count"] == 1
    assert len(result["semantic_missed_samples"]) == 1
    text = extract_sample_text(out_path)
    assert "内控审计项目" in text
    assert "不存在的内容" not in text


def test_semantic_apply_hits_after_entity_replace(tmp_path):
    doc_path = tmp_path / "combo.docx"
    _minimal_docx(doc_path, ["中国能建251231内控审计"])
    out_path = tmp_path / "out.docx"
    entities = [
        {
            "canonical_name": "中国能建",
            "entity_type": "company",
            "placeholder": "[公司_1]",
            "source": "manual",
            "is_excluded": False,
            "aliases": ["中国能建"],
        },
    ]
    pairs = [
        {
            "original": "251231内控审计",
            "rewritten": "内控审计项目",
            "category": "project_id",
        },
    ]
    result = run_deid_pipeline(
        doc_path, out_path, entities, [], [], semantic_pairs=pairs
    )
    assert result["semantic_applied_count"] == 1
    assert result["semantic_missed_count"] == 0
    text = extract_sample_text(out_path)
    assert "[公司_1]" in text
    assert "内控审计项目" in text
    assert "251231" not in text


def test_partition_semantic_pairs_on_unpacked_workdir(tmp_path):
    work = tmp_path / "work"
    work.mkdir()
    word = work / "word"
    word.mkdir()
    doc_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/document/2006/main">'
        "<w:body><w:p><w:r><w:t>命中段落</w:t></w:r></w:p>"
        "</w:body></w:document>"
    )
    (word / "document.xml").write_text(doc_xml, encoding="utf-8")
    applicable, missed = partition_semantic_pairs(
        work,
        [
            {"original": "命中段落", "rewritten": "替换"},
            {"original": "缺失", "rewritten": "x"},
        ],
    )
    assert len(applicable) == 1
    assert len(missed) == 1
